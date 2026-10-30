"""Recompute the headline numbers of the report from the raw result files and compare them with the values printed in the report.

Usage (from the repository root or from the unpacked Evaluation folder, CPU only, about 10 seconds):
    python tools/verify_report_numbers.py                       recompute and compare with the reported values
    python tools/verify_report_numbers.py --report "49030000 Report.pdf"      also check that each reported phrase is in the report (.pdf or .docx)

Every row reads: id, what is checked, the recomputed value, the value printed in the report, PASS or FAIL.
A recomputed value passes when it equals the reported value after rounding to the reported number of decimals.
With --report, a second column says whether the phrase that carries the number was found in the report text.
Exit code 1 if any row fails. The raw files used are listed in the "source" column; nothing is written.
"""
import argparse, json, pathlib, re, string, sys
import numpy as np
import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent.parent
P2, P4, P6, P7 = (ROOT / "pairs" / n / "results" for n in ("p2_3d_reconstruction", "p4_agents", "p6_lmm_vqa", "p7_safety"))
ROWS = []   # (id, what, recomputed, reported, decimals, phrase, source)


def jl(path):
    return pd.DataFrame([json.loads(l) for l in open(path, encoding="utf-8") if l.strip()])


def add(i, what, value, reported, decimals, phrase, source):
    ROWS.append((i, what, float(value), float(reported), decimals, phrase, source))


# ------------------------------------------------------------------ pair 7: pooled over both runs, a prompt drawn in both runs counted once
P = pd.concat([jl(P7 / "predictions.jsonl"), jl(P7 / "replicate_seed1" / "predictions.jsonl")]).drop_duplicates(["set", "uid", "model"])
def p7(set_, model, label):
    g = P[(P.set == set_) & (P.model == model) & (P.label == label)]; return 100 * g.decision.mean()
A, B = "llama_guard_3_1b", "shieldgemma_2b"
add("7.1", "ToxicChat recall, Llama Guard (%)", p7("toxicchat", A, 1), 57.1, 1, "from 57.1% to 71.0%", "pair 7 predictions, both runs")
add("7.2", "ToxicChat recall, ShieldGemma (%)", p7("toxicchat", B, 1), 71.0, 1, "from 57.1% to 71.0%", "pair 7 predictions, both runs")
add("7.3", "ToxicChat false blocks, Llama Guard (%)", p7("toxicchat", A, 0), 14.4, 1, "from 14.4% to 2.9%", "pair 7 predictions, both runs")
add("7.4", "ToxicChat false blocks, ShieldGemma (%)", p7("toxicchat", B, 0), 2.9, 1, "from 14.4% to 2.9%", "pair 7 predictions, both runs")
add("7.5", "HateCheck false blocks, Llama Guard (%)", p7("hatecheck", A, 0), 17.7, 1, "17.7% to 46.6%", "pair 7 predictions, both runs")
add("7.6", "HateCheck false blocks, ShieldGemma (%)", p7("hatecheck", B, 0), 46.6, 1, "17.7% to 46.6%", "pair 7 predictions, both runs")
cf = pd.read_csv(P7 / "extra" / "combined_filters.csv")
r = cf[(cf.set == "toxicchat") & (cf.rule == "either flags")].iloc[0]
add("7.7", "ToxicChat recall, either filter flags (%)", 100 * r.recall, 82.8, 1, "82.8%", "pair 7 extra/combined_filters.csv")
add("7.8", "ToxicChat false blocks, either filter flags (%)", 100 * r.false_block_rate, 16.5, 1, "16.5%", "pair 7 extra/combined_filters.csv")
g = pd.read_csv(P7 / "extra" / "hatecheck_by_group.csv")
def grp(model, group, measure): return 100 * g[(g.model == model) & (g.group == group) & (g.measure == measure)].rate.iloc[0]
add("7.9", "HateCheck recall, Llama Guard, women (%)", grp(A, "women", "recall on hateful"), 64, 0, "64% of the hateful cases that target women", "pair 7 extra/hatecheck_by_group.csv")
add("7.10", "HateCheck recall, Llama Guard, disabled people (%)", grp(A, "disabled people", "recall on hateful"), 76, 0, "76% for disabled people", "pair 7 extra/hatecheck_by_group.csv")
add("7.11", "HateCheck false blocks, ShieldGemma, Black people (%)", grp(B, "black people", "false-block rate on non-hateful"), 64, 0, "64% of the harmless cases that name Black people", "pair 7 extra/hatecheck_by_group.csv")
add("7.12", "HateCheck false blocks, ShieldGemma, Muslims (%)", grp(B, "Muslims", "false-block rate on non-hateful"), 61, 0, "61% of those that name Muslims", "pair 7 extra/hatecheck_by_group.csv")
sm = pd.read_csv(P7 / "summary.csv")
add("7.13", "ShieldGemma AUC on ToxicChat, run 1", sm[(sm.system.str.startswith("B")) & (sm.dataset == "toxicchat")].auc.iloc[0], 0.96, 2, "AUC on this set is 0.96", "pair 7 summary.csv")

# ------------------------------------------------------------------ pair 6
def yn(s):
    w = str(s).lower().strip().rstrip(".").translate(str.maketrans("", "", string.punctuation)).split()
    return w[0] if w and w[0] in ("yes", "no") else "invalid"
man = {x["id"]: x["truth"] for f in (P6.parent / "data" / "eval_manifest.json", P6.parent / "data_seed1" / "eval_manifest.json") for x in json.load(open(f))}
pr = pd.concat([jl(P6 / "predictions.jsonl"), jl(P6 / "replicate_seed1" / "predictions.jsonl")])
pr = pr[(pr.set == "pope") & (pr.variant == "original")].drop_duplicates(["model", "id"]); pr["y"] = pr.prediction.map(yn); pr["t"] = pr.id.map(man)
def pool(model, kind):
    g = pr[pr.model == model]
    return 100 * ((g.y == "yes") & (g.t == "no")).sum() / (g.t == "no").sum() if kind == "halluc" else 100 * ((g.y == "yes") & (g.t == "yes")).sum() / (g.t == "yes").sum()
QA, IB, QR = "qwen3-vl-8b-instruct (ollama q4)", "internvl3-8b (bnb nf4)", "qwen2.5-vl-7b (ollama q4)"
add("6.1", "Pooled hallucination rate, Qwen3-VL (%)", pool(QA, "halluc"), 7.8, 1, "7.8%", "pair 6 predictions, both samples")
add("6.2", "Pooled hallucination rate, InternVL3 (%)", pool(IB, "halluc"), 13.5, 1, "13.5%", "pair 6 predictions, both samples")
add("6.3", "Pooled hallucination rate, Qwen2.5-VL reference (%)", pool(QR, "halluc"), 2.1, 1, "2.1%", "pair 6 predictions, both samples")
add("6.4", "Pooled recall on present objects, Qwen3-VL (%)", pool(QA, "recall"), 82.2, 1, "82.2% against 87.5% recall", "pair 6 predictions, both samples")
add("6.5", "Pooled recall on present objects, InternVL3 (%)", pool(IB, "recall"), 87.5, 1, "82.2% against 87.5% recall", "pair 6 predictions, both samples")
inj = jl(P6 / "extra" / "injection_predictions.jsonl"); inj["y"] = inj.prediction.map(yn)
absent = {i for i, t in man.items() if t == "no" and i.startswith("pope")}
cl = pr[pr.id.isin(set(inj.id)) & (pr.set == "pope")]
yes = lambda df: int((df.y == "yes").sum())
a_inj = inj[inj.model == QA]; b_inj = inj[inj.model == IB]
add("6.6", "Image-text injection: Qwen3-VL said yes, injected banner (of 150)", yes(a_inj[a_inj.cond == "injected"]), 55, 0, "55 (36.7%)", "pair 6 extra/injection_predictions.jsonl")
add("6.7", "Image-text injection: Qwen3-VL said yes, no banner (of 150)", yes(cl[cl.model == QA]), 16, 0, "16 cases without a banner", "pair 6 predictions.jsonl (same questions)")
add("6.8", "Image-text injection: InternVL3 said yes, injected banner (of 150)", yes(b_inj[b_inj.cond == "injected"]), 25, 0, "21, 21 and 25 cases", "pair 6 extra/injection_predictions.jsonl")

# ------------------------------------------------------------------ pair 4: pooled over seeds 0 to 19
E = pd.concat([jl(P4 / "episodes.jsonl"), jl(P4 / "replicate_seed1" / "episodes.jsonl")])
m = E[E.phase == "main"].groupby("agent").success.mean() * 100
add("4.1", "Pooled success, text agent (%)", m["text_qwen2.5-7b"], 79.5, 1, "79.5% against 72.5%", "pair 4 episodes, both runs")
add("4.2", "Pooled success, vision agent (%)", m["vlm_qwen2.5vl-7b"], 72.5, 1, "79.5% against 72.5%", "pair 4 episodes, both runs")
ij = E[E.phase == "injection"].groupby("agent").attack_followed.sum()
add("4.3", "Injected instruction followed, text agent (of 60)", ij["text_qwen2.5-7b"], 30, 0, "30 of 60", "pair 4 episodes, both runs")
add("4.4", "Injected instruction followed, vision agent (of 60)", ij["vlm_qwen2.5vl-7b"], 5, 0, "5 of 60", "pair 4 episodes, both runs")

# ------------------------------------------------------------------ pair 2: mean of the two seeds
pm = pd.read_csv(P2 / "pooled" / "p2_mean_and_range.csv").set_index(["scene", "method"]).psnr_mean
add("2.1", "PSNR, splatfacto, poster scene, mean of two seeds (dB)", pm[("poster", "splatfacto")], 32.9, 1, "32.9 against 20.5 dB", "pair 2 pooled/p2_mean_and_range.csv")
add("2.2", "PSNR, nerfacto, poster scene, mean of two seeds (dB)", pm[("poster", "nerfacto")], 20.5, 1, "32.9 against 20.5 dB", "pair 2 pooled/p2_mean_and_range.csv")
add("2.3", "PSNR, splatfacto, chair video, mean of two seeds (dB)", pm[("chair_video", "splatfacto")], 24.7, 1, "24.7 against 19.8 dB", "pair 2 pooled/p2_mean_and_range.csv")
add("2.4", "PSNR, nerfacto, chair video, mean of two seeds (dB)", pm[("chair_video", "nerfacto")], 19.8, 1, "24.7 against 19.8 dB", "pair 2 pooled/p2_mean_and_range.csv")


def report_text(path):
    path = pathlib.Path(path)
    if path.suffix.lower() == ".pdf":
        import pypdf
        t = "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(path)).pages)
    else:
        import docx
        d = docx.Document(str(path)); t = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for tb in d.tables for r in tb.rows for c in r.cells)
    return re.sub(r"\s+", " ", t)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--report"); a = ap.parse_args()
    text = report_text(a.report) if a.report else None
    bad = 0
    print(f"{'id':5} {'what':62} {'recomputed':>11} {'reported':>9}  value  " + ("phrase  " if text else "") + "source")
    for i, what, val, rep, dec, phrase, src in ROWS:
        ok = round(val, dec) == round(rep, dec); found = None if text is None else (re.sub(r"\s+", " ", phrase) in text)
        bad += (not ok) or (found is False)
        print(f"{i:5} {what[:62]:62} {val:11.{dec + 1}f} {rep:9.{dec}f}  {'PASS' if ok else 'FAIL':5}  " + ("" if found is None else f"{'found' if found else 'MISSING':8}") + src)
    print(f"\n{len(ROWS)} numbers checked, {bad} problem(s)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
